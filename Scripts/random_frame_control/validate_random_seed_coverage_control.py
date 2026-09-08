#!/usr/bin/env python3
"""Independent replay and share-package validation for the coverage control."""

from __future__ import annotations

import csv
import gzip
import hashlib
import itertools
import json
import math
import os
import platform
import sys
from collections import Counter
from pathlib import Path

import nbformat
import numpy as np
import openpyxl
import scipy
from PIL import Image
from scipy.spatial.distance import cdist


PACKAGE = Path(__file__).resolve().parents[1]
INPUTS = PACKAGE / "inputs"
RESULTS = PACKAGE / "results"
TRAJECTORIES = PACKAGE / "trajectories"
TYPE_MAP = ("Hf", "Mo", "Ta", "Ti", "Zr")
RADIAL_CENTERS = np.arange(1.50, 6.00, 0.15, dtype=np.float64)
SIGMA = 0.15
CUTOFF = 6.0
MASTER_SEED = 20260903
N_DRAWS = 10_000
FOLD_CODES = {"fcc_to_bcc": 1, "bcc_to_fcc": 2}
CONTROL_CODES = {"pooled_random_frames": 1, "one_per_trajectory": 2}
EXPECTED_SHEETS = [
    "README",
    "Figure",
    "Control_summary",
    "Draw_distribution",
    "System_summary",
    "Candidate_frames",
    "Trajectories",
    "Frames_and_CIFs",
    "Seed_reconstruction",
    "Frozen_contrasts",
    "Source_hashes",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def read_gzip_csv(path: Path) -> list[dict[str, str]]:
    with gzip.open(path, "rt", newline="") as handle:
        return list(csv.DictReader(handle))


def load_deepmd(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    if tuple((path / "type_map.raw").read_text().split()) != TYPE_MAP:
        raise AssertionError(f"Type-map failure: {path}")
    atom_types = np.asarray(np.loadtxt(path / "type.raw", dtype=int, ndmin=1))
    coordinates = []
    boxes = []
    for set_dir in sorted(path.glob("set.*")):
        coordinates.append(np.load(set_dir / "coord.npy", allow_pickle=False))
        boxes.append(np.load(set_dir / "box.npy", allow_pickle=False))
    coords = np.concatenate(coordinates).reshape(-1, len(atom_types), 3)
    cells = np.concatenate(boxes).reshape(-1, 3, 3)
    return coords, cells, atom_types


def selected_atoms(atom_types: np.ndarray) -> np.ndarray:
    selected: list[int] = []
    for element_index in range(5):
        eligible = np.flatnonzero(atom_types == element_index)
        if len(eligible) <= 12:
            selected.extend(eligible.tolist())
        else:
            positions = np.rint(np.linspace(0, len(eligible) - 1, 12)).astype(int)
            selected.extend(eligible[positions].tolist())
    return np.asarray(selected, dtype=int)


def independent_frame_descriptor(
    coordinates: np.ndarray, cell: np.ndarray, atom_types: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Separate raw-coordinate implementation used only by the validator."""

    inverse = np.linalg.inv(cell)
    fractional = coordinates @ inverse
    plane_heights = 1.0 / np.linalg.norm(inverse, axis=0)
    spans = np.ceil(CUTOFF / plane_heights).astype(int)
    shifts = np.asarray(
        list(
            itertools.product(
                range(-int(spans[0]), int(spans[0]) + 1),
                range(-int(spans[1]), int(spans[1]) + 1),
                range(-int(spans[2]), int(spans[2]) + 1),
            )
        ),
        dtype=float,
    )
    atom_indices = selected_atoms(atom_types)
    descriptors = np.zeros((len(atom_indices), 5, len(RADIAL_CENTERS)), dtype=float)
    for output_position, central_index in enumerate(atom_indices):
        displacement = fractional - fractional[central_index]
        displacement -= np.floor(displacement + 0.5)
        repeated = displacement[:, None, :] + shifts[None, :, :]
        distances = np.linalg.norm(repeated @ cell, axis=2).reshape(-1)
        neighbour_types = np.repeat(atom_types, len(shifts))
        use = (distances > 1.0e-10) & (distances < CUTOFF)
        retained = distances[use]
        taper = 0.5 * (np.cos(np.pi * retained / CUTOFF) + 1.0)
        radial = np.exp(
            -0.5 * ((retained[:, None] - RADIAL_CENTERS[None, :]) / SIGMA) ** 2
        )
        radial *= taper[:, None]
        for neighbour_type, values in zip(neighbour_types[use], radial, strict=True):
            descriptors[output_position, int(neighbour_type)] += values
    return atom_indices, atom_types[atom_indices], descriptors.reshape(len(atom_indices), -1)


def expected_selections(
    fold: str, control: str, candidate_rows: list[dict[str, str]]
) -> np.ndarray:
    system_codes = np.asarray([int(row["candidate_system_code"]) for row in candidate_rows])
    number_of_systems = len(np.unique(system_codes))
    rng = np.random.Generator(
        np.random.PCG64(
            np.random.SeedSequence([MASTER_SEED, FOLD_CODES[fold], CONTROL_CODES[control]])
        )
    )
    if control == "pooled_random_frames":
        output = np.empty((N_DRAWS, number_of_systems), dtype=np.int16)
        for draw_index in range(N_DRAWS):
            output[draw_index] = rng.choice(
                len(candidate_rows), size=number_of_systems, replace=False
            )
        return output
    columns = []
    for code in range(number_of_systems):
        eligible = np.flatnonzero(system_codes == code)
        columns.append(rng.choice(eligible, size=N_DRAWS, replace=True))
    return np.column_stack(columns).astype(np.int16)


def system_slices(
    fold: str, design: dict, descriptor_manifest: dict[str, dict[str, str]]
) -> tuple[list[str], list[tuple[int, int]]]:
    test_ids = sorted(design["folds"][fold]["pocc_test_source_ids"])
    slices = []
    cursor = 0
    for system_id in test_ids:
        count = int(descriptor_manifest[system_id]["environment_count"])
        slices.append((cursor, cursor + count))
        cursor += count
    return test_ids, slices


def replay_fold_statistics(
    matrix: np.ndarray,
    baseline: np.ndarray,
    selections: np.ndarray,
    slices: list[tuple[int, int]],
) -> np.ndarray:
    pure_medians = np.asarray([np.median(baseline[start:stop]) for start, stop in slices])
    output = np.empty(len(selections), dtype=float)
    for batch_start in range(0, len(selections), 41):
        batch_stop = min(batch_start + 41, len(selections))
        nearest_random = np.take(matrix, selections[batch_start:batch_stop], axis=0).min(axis=1)
        augmented = np.minimum(nearest_random, baseline[None, :])
        reductions = []
        for system_index, (start, stop) in enumerate(slices):
            medians = np.median(augmented[:, start:stop], axis=1)
            reductions.append((pure_medians[system_index] - medians) / pure_medians[system_index])
        output[batch_start:batch_stop] = np.median(np.column_stack(reductions), axis=1)
    return output


def validate_cif(path: Path, natoms: int) -> None:
    lines = path.read_text().splitlines()
    occupancy_header = lines.index("_atom_site_occupancy")
    atom_lines = [line for line in lines[occupancy_header + 1 :] if line.strip()]
    if len(atom_lines) != natoms:
        raise AssertionError(f"CIF atom count mismatch: {path}")
    for line in atom_lines:
        fields = line.split()
        if len(fields) != 6 or fields[1] not in TYPE_MAP or float(fields[5]) != 1.0:
            raise AssertionError(f"Malformed CIF atom row: {path}")
        fractional = [float(value) for value in fields[2:5]]
        if any(not (0.0 <= value < 1.0 + 1.0e-12) for value in fractional):
            raise AssertionError(f"CIF fractional coordinate outside cell: {path}")


def main() -> int:
    checks = 0

    def checked(condition: bool, message: str) -> None:
        nonlocal checks
        if not condition:
            raise AssertionError(message)
        checks += 1

    source_rows = read_csv(INPUTS / "source_file_manifest.csv")
    checked(len(source_rows) >= 800, "Source manifest unexpectedly small")
    for row in source_rows:
        copied = PACKAGE / row["copied_path"]
        checked(copied.is_file(), f"Missing copied source: {copied}")
        checked(copied.stat().st_size == int(row["bytes"]), f"Size mismatch: {copied}")
        checked(sha256(copied) == row["sha256"], f"Hash mismatch: {copied}")
        source = Path(row["source_path"])
        checked(source.is_file(), f"Missing immutable source: {source}")
        checked(sha256(source) == row["sha256"], f"Source drift: {source}")

    trajectories = read_csv(RESULTS / "trajectory_manifest.csv")
    frames = read_csv(RESULTS / "frame_manifest.csv")
    candidates = read_csv(RESULTS / "candidate_frames.csv")
    controls = read_csv(RESULTS / "control_summary.csv")
    draws = read_gzip_csv(RESULTS / "random_draw_distribution.csv.gz")
    reconstruction = read_csv(RESULTS / "systematic_seed_reconstruction.csv")
    checked(len(trajectories) == 103, "Expected 103 analyzed trajectories")
    checked(sum(row["family"] == "PURE" for row in trajectories) == 20, "PURE count")
    checked(sum(row["family"] == "POCC" for row in trajectories) == 83, "POCC count")
    checked(len(frames) == 577, "Expected 577 analyzed frames")
    checked(len(candidates) == 415, "Expected 415 candidate frames")
    checked(len(draws) == 40_000, "Expected 40,000 draw rows")
    checked(len(controls) == 4, "Expected four control summaries")

    for row in trajectories:
        copied = PACKAGE / row["deepmd_path"]
        checked(copied.is_dir(), f"Missing trajectory directory: {copied}")
        checked(
            len(read_csv(copied / "frame_selection.csv"))
            == int(row["analyzed_frames_temporally_thinned"]),
            f"Trajectory frame count: {row['system_id']}",
        )
    for row in frames:
        cif_path = PACKAGE / row["cif_path"]
        checked(cif_path.is_file(), f"Missing CIF: {cif_path}")
        checked(sha256(cif_path) == row["cif_sha256"], f"CIF hash: {cif_path}")
        validate_cif(cif_path, int(row["natoms"]))
        checks += 1

    design = json.loads((INPUTS / "study_design.json").read_text())
    descriptor_rows = read_csv(INPUTS / "descriptor_cache_manifest.csv")
    descriptor_manifest = {row["system_id"]: row for row in descriptor_rows}
    checked(len(descriptor_manifest) == 103, "Descriptor manifest count")

    # Raw-coordinate descriptor replay for two frames in one PURE, FCC, and BCC system.
    selected_systems = [
        next(row["system_id"] for row in trajectories if row["family"] == "PURE"),
        next(
            row["system_id"]
            for row in trajectories
            if row["family"] == "POCC" and row["parent_family"] == "fcc"
        ),
        next(
            row["system_id"]
            for row in trajectories
            if row["family"] == "POCC" and row["parent_family"] == "bcc"
        ),
    ]
    trajectory_by_id = {row["system_id"]: row for row in trajectories}
    for system_id in selected_systems:
        coords, cells, atom_types = load_deepmd(PACKAGE / trajectory_by_id[system_id]["deepmd_path"])
        cache_path = PACKAGE / descriptor_manifest[system_id]["cache_path"]
        with np.load(cache_path, allow_pickle=False) as cache:
            cached = {name: cache[name] for name in cache.files}
        for frame_index in sorted({0, len(coords) - 1}):
            atom_indices, elements, descriptors = independent_frame_descriptor(
                coords[frame_index], cells[frame_index], atom_types
            )
            mask = cached["frame_index"] == frame_index
            checked(np.array_equal(atom_indices, cached["atom_index"][mask]), "Atom selection replay")
            checked(np.array_equal(elements, cached["element_index"][mask]), "Element replay")
            checked(
                np.allclose(descriptors, cached["descriptors"][mask], rtol=2.0e-6, atol=2.0e-6),
                f"Raw descriptor replay: {system_id}/{frame_index}",
            )

    candidate_by_fold = {
        fold: sorted(
            [row for row in candidates if row["fold"] == fold],
            key=lambda row: int(row["candidate_index"]),
        )
        for fold in FOLD_CODES
    }
    checked(len(candidate_by_fold["fcc_to_bcc"]) == 175, "FCC candidate count")
    checked(len(candidate_by_fold["bcc_to_fcc"]) == 240, "BCC candidate count")
    checked(
        sum(row["systematic_first_frame"] == "True" for row in candidate_by_fold["fcc_to_bcc"])
        == 29,
        "FCC systematic count",
    )
    checked(
        sum(row["systematic_first_frame"] == "True" for row in candidate_by_fold["bcc_to_fcc"])
        == 30,
        "BCC systematic count",
    )

    with np.load(RESULTS / "random_selections.npz", allow_pickle=False) as stored:
        stored_selections = {name: stored[name] for name in stored.files}

    draw_lookup: dict[tuple[str, str], np.ndarray] = {}
    for fold in FOLD_CODES:
        for control in CONTROL_CODES:
            values = sorted(
                (
                    (int(row["draw_index"]), float(row["coverage_reduction_fraction"]))
                    for row in draws
                    if row["fold"] == fold and row["control"] == control
                ),
                key=lambda pair: pair[0],
            )
            checked(len(values) == N_DRAWS, f"Draw count {fold}/{control}")
            checked([index for index, _ in values] == list(range(N_DRAWS)), "Draw indices")
            draw_lookup[(fold, control)] = np.asarray([value for _, value in values])

    control_lookup = {(row["fold"], row["control"]): row for row in controls}
    direct_distance_rows_checked = 0
    for fold in FOLD_CODES:
        cache_path = RESULTS / "distance_cache" / f"{fold}_candidate_to_test_distances.npz"
        with np.load(cache_path, allow_pickle=False) as cache:
            matrix = cache["candidate_distance"]
            baseline = cache["baseline_distance"]
            frozen_seed = cache["frozen_seed_distance"]
        test_ids, slices = system_slices(fold, design, descriptor_manifest)
        checked(matrix.shape[1] == slices[-1][1], f"Test environment shape {fold}")
        checked(np.isfinite(matrix).all() and np.all(matrix >= 0), f"Distance matrix finite {fold}")

        test_descriptors = []
        test_elements = []
        for system_id in test_ids:
            with np.load(PACKAGE / descriptor_manifest[system_id]["cache_path"], allow_pickle=False) as data:
                test_descriptors.append(data["descriptors"].astype(float))
                test_elements.append(data["element_index"])
        test_descriptor = np.concatenate(test_descriptors)
        test_element = np.concatenate(test_elements)

        systematic_positions = [
            int(row["candidate_index"])
            for row in candidate_by_fold[fold]
            if row["systematic_first_frame"] == "True"
        ]
        extra_positions = np.linspace(0, len(candidate_by_fold[fold]) - 1, 12).round().astype(int).tolist()
        for candidate_position in sorted(set(systematic_positions + extra_positions)):
            candidate = candidate_by_fold[fold][candidate_position]
            with np.load(
                PACKAGE / descriptor_manifest[candidate["system_id"]]["cache_path"],
                allow_pickle=False,
            ) as data:
                frame_mask = data["frame_index"] == int(candidate["canonical_frame_index"])
                candidate_descriptor = data["descriptors"][frame_mask].astype(float)
                candidate_element = data["element_index"][frame_mask]
            replay = np.empty(len(test_descriptor), dtype=float)
            for element_index in range(5):
                query = test_element == element_index
                reference = candidate_element == element_index
                replay[query] = cdist(
                    test_descriptor[query], candidate_descriptor[reference], metric="euclidean"
                ).min(axis=1)
            checked(
                np.allclose(replay, matrix[candidate_position], rtol=2.0e-12, atol=2.0e-12),
                f"Direct candidate-distance replay {fold}/{candidate_position}",
            )
            direct_distance_rows_checked += 1

        systematic_augmented = np.minimum(baseline, matrix[systematic_positions].min(axis=0))
        checked(
            np.allclose(systematic_augmented, frozen_seed, rtol=0.0, atol=1.0e-10),
            f"Frozen seed environment replay {fold}",
        )
        system_reductions = []
        for start, stop in slices:
            pure_median = np.median(baseline[start:stop])
            seed_median = np.median(systematic_augmented[start:stop])
            system_reductions.append((pure_median - seed_median) / pure_median)
        observed = float(np.median(system_reductions))

        for control in CONTROL_CODES:
            expected = expected_selections(fold, control, candidate_by_fold[fold])
            saved = stored_selections[f"{fold}__{control}"]
            checked(np.array_equal(expected, saved), f"RNG replay {fold}/{control}")
            replayed = replay_fold_statistics(matrix, baseline, expected, slices)
            checked(
                np.allclose(replayed, draw_lookup[(fold, control)], rtol=0.0, atol=1.0e-14),
                f"All draw statistics replay {fold}/{control}",
            )
            summary = control_lookup[(fold, control)]
            checked(
                math.isclose(float(summary["systematic_seed_reduction_fraction"]), observed, abs_tol=1e-14),
                f"Observed summary {fold}/{control}",
            )
            checked(
                math.isclose(float(summary["random_median_reduction_fraction"]), np.median(replayed), abs_tol=1e-14),
                f"Random median {fold}/{control}",
            )
            checked(
                math.isclose(float(summary["random_q025_reduction_fraction"]), np.quantile(replayed, 0.025), abs_tol=1e-14),
                f"Random q025 {fold}/{control}",
            )
            checked(
                math.isclose(float(summary["random_q975_reduction_fraction"]), np.quantile(replayed, 0.975), abs_tol=1e-14),
                f"Random q975 {fold}/{control}",
            )
            checked(np.all(replayed > observed), f"Expected every random draw to exceed observed: {fold}/{control}")

    checked(all(row["status"] == "passed" for row in reconstruction), "Seed reconstruction status")
    checked(
        all(float(row["maximum_environment_distance_difference"]) <= 1.0e-10 for row in reconstruction),
        "Seed reconstruction tolerance",
    )

    png_path = PACKAGE / "figures" / "random_seed_coverage_control.png"
    svg_path = PACKAGE / "figures" / "random_seed_coverage_control.svg"
    with Image.open(png_path) as image:
        checked(image.width >= 3000 and image.height >= 1200, "Figure pixel dimensions")
        dpi = image.info.get("dpi", (0, 0))
        checked(min(dpi) >= 299.0, "Figure DPI")
    svg_text = svg_path.read_text()
    checked("<svg" in svg_text and "<image" not in svg_text, "Editable vector figure")

    workbook_path = PACKAGE / "Random_Seed_Coverage_Control.xlsx"
    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=False)
    checked(workbook.sheetnames == EXPECTED_SHEETS, "Workbook sheets")
    expected_rows = {
        "Control_summary": 4,
        "Draw_distribution": 40_000,
        "System_summary": 166,
        "Candidate_frames": 415,
        "Trajectories": 103,
        "Frames_and_CIFs": 577,
        "Seed_reconstruction": 2,
        "Frozen_contrasts": 166,
        "Source_hashes": len(source_rows),
    }
    for sheet, count in expected_rows.items():
        checked(workbook[sheet].max_row == count + 1, f"Workbook row count: {sheet}")
    checked(len(workbook["Figure"]._images) == 1, "Workbook embedded figure")
    workbook.close()

    notebook_path = PACKAGE / "notebooks" / "random_seed_coverage_control.ipynb"
    notebook = nbformat.read(notebook_path, as_version=4)
    code_cells = [cell for cell in notebook.cells if cell.cell_type == "code"]
    checked(len(code_cells) >= 3, "Notebook code-cell count")
    checked(all(cell.execution_count is not None for cell in code_cells), "Notebook executed")
    checked(
        not any(
            output.output_type == "error"
            for cell in code_cells
            for output in cell.get("outputs", [])
        ),
        "Notebook error outputs",
    )

    receipt = {
        "schema_version": 1,
        "status": "passed",
        "checks": checks,
        "direct_candidate_distance_rows_replayed": direct_distance_rows_checked,
        "random_draw_statistics_replayed": 40_000,
        "raw_descriptor_frames_replayed": 6,
        "cif_files_validated": len(frames),
        "source_manifest_rows_validated": len(source_rows),
        "trajectory_count": len(trajectories),
        "analyzed_frame_count": len(frames),
        "candidate_frame_count": len(candidates),
        "environment": {
            "python": sys.version,
            "platform": platform.platform(),
            "numpy": np.__version__,
            "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
            "nbformat": nbformat.__version__,
        },
    }
    path = RESULTS / "validation_receipt.json"
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temporary.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n")
    temporary.replace(path)
    print(json.dumps(receipt, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"VALIDATION ERROR: {error}", file=sys.stderr)
        raise
