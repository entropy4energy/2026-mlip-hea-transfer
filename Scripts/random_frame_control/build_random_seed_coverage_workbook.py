#!/usr/bin/env python3
"""Build the colleague-facing Excel workbook for the random-frame control."""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path

import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PACKAGE = Path(__file__).resolve().parents[1]
RESULTS = PACKAGE / "results"
INPUTS = PACKAGE / "inputs"
FIGURES = PACKAGE / "figures"
OUTPUT = PACKAGE / "Random_Seed_Coverage_Control.xlsx"
INK = "263238"
BLUE = "31688E"
GOLD = "D9A441"
PALE_BLUE = "EAF1F5"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_tables() -> dict[str, pd.DataFrame]:
    tables = {
        "Control_summary": pd.read_csv(RESULTS / "control_summary.csv"),
        "Draw_distribution": pd.read_csv(RESULTS / "random_draw_distribution.csv.gz"),
        "System_summary": pd.read_csv(RESULTS / "system_random_control_summary.csv"),
        "Candidate_frames": pd.read_csv(RESULTS / "candidate_frames.csv"),
        "Trajectories": pd.read_csv(RESULTS / "trajectory_manifest.csv"),
        "Frames_and_CIFs": pd.read_csv(RESULTS / "frame_manifest.csv"),
        "Seed_reconstruction": pd.read_csv(RESULTS / "systematic_seed_reconstruction.csv"),
        "Frozen_contrasts": pd.read_csv(INPUTS / "frozen_environment_coverage_contrasts.csv"),
        "Source_hashes": pd.read_csv(INPUTS / "source_file_manifest.csv"),
    }
    return tables


def style_data_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 32
    for column_index, cells in enumerate(worksheet.iter_cols(min_row=1), start=1):
        header = str(cells[0].value or "")
        sample_lengths = [len(str(cell.value)) for cell in cells[: min(len(cells), 300)] if cell.value is not None]
        width = min(max([len(header) + 2, *sample_lengths], default=12) + 1, 58)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 11)
        if any(token in header for token in ("fraction", "p_random", "percentile")):
            number_format = "0.0000"
        elif "percent" in header or "percentage_points" in header:
            number_format = "0.00"
        elif any(token in header for token in ("distance", "energy", "cell_", "volume")):
            number_format = "0.000000"
        else:
            number_format = None
        if number_format:
            for cell in cells[1:]:
                cell.number_format = number_format


def main() -> int:
    tables = read_tables()
    temporary = OUTPUT.with_name(f".{OUTPUT.name}.{os.getpid()}.tmp.xlsx")
    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        for sheet_name, dataframe in tables.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
        workbook = writer.book
        readme = workbook.create_sheet("README", 0)
        readme.sheet_view.showGridLines = False
        readme.column_dimensions["A"].width = 25
        readme.column_dimensions["B"].width = 110
        rows = [
            ("Random-seed coverage control", "Matched-label reference requested by the PI; internal scientific-review data workbook."),
            ("Analysis date", "2026-09-03"),
            ("Frozen representation", "150-dimensional label-blind chemical radial descriptor; same-central-element Euclidean distance."),
            ("Frame population", "Frozen every-100th-frame temporal-thinning variants; reduced serial dependence is not statistical independence."),
            ("Primary control", "10,000 uniform draws without replacement of 29 FCC-parent or 30 BCC-parent frames from the eligible pooled candidate frames."),
            ("Sensitivity", "10,000 draws of exactly one random frame from each eligible training-parent trajectory."),
            ("Headline estimand", "Median across held-out systems of each system's paired relative reduction in median nearest-training distance from PURE."),
            ("Important boundary", "Random controls use POCC-derived trajectory frames. They do not compare POCC with random-alloy, SQS, or active-learning structure generation, and no models were trained."),
            ("CIF data", "See the Frames_and_CIFs sheet for one row and SHA-256 per CIF; files are under trajectories/cif/."),
            ("DeepMD trajectories", "Exact analyzed trajectory directories are under trajectories/deepmd/ and indexed in the Trajectories and Source_hashes sheets."),
            ("Control_summary", "Four fold × control headline rows, including systematic values, random distributions, empirical percentile, and tail probability."),
            ("Draw_distribution", "All 40,000 fold-level random-draw statistics."),
            ("System_summary", "Held-out-system random-control summaries."),
            ("Candidate_frames", "All 415 eligible candidate frames and their source/CIF provenance."),
            ("Seed_reconstruction", "Exact reproduction check against the frozen first-frame coverage result."),
            ("Source_hashes", "File-level hashes tying the copied inputs and trajectories to their immutable source files."),
        ]
        for row_index, (label, value) in enumerate(rows, start=1):
            readme.cell(row_index, 1, label)
            readme.cell(row_index, 2, value)
            readme.cell(row_index, 1).font = Font(bold=True, color="FFFFFF" if row_index == 1 else INK)
            readme.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
            readme.cell(row_index, 1).alignment = Alignment(wrap_text=True, vertical="top")
            if row_index == 1:
                readme.cell(row_index, 1).fill = PatternFill("solid", fgColor=INK)
                readme.cell(row_index, 2).fill = PatternFill("solid", fgColor=INK)
                readme.cell(row_index, 2).font = Font(color="FFFFFF", bold=True)
            elif row_index % 2 == 0:
                readme.cell(row_index, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
                readme.cell(row_index, 2).fill = PatternFill("solid", fgColor=PALE_BLUE)
            readme.row_dimensions[row_index].height = 30 if row_index > 1 else 40

        for name in tables:
            style_data_sheet(workbook[name])

        figure_sheet = workbook.create_sheet("Figure", 1)
        figure_sheet.sheet_view.showGridLines = False
        figure_sheet["A1"] = "Matched-label random-frame coverage control"
        figure_sheet["A1"].font = Font(size=16, bold=True, color=INK)
        figure_sheet["A2"] = (
            "Violin distributions contain 10,000 draws per fold and control; see Control_summary and Draw_distribution for exact values."
        )
        figure_sheet["A2"].alignment = Alignment(wrap_text=True)
        figure_sheet.column_dimensions["A"].width = 110
        image = Image(FIGURES / "random_seed_coverage_control.png")
        image.width = 1054
        image.height = 459
        figure_sheet.add_image(image, "A4")

        workbook.properties.title = "Random-seed coverage control"
        workbook.properties.subject = "Matched-label structural-coverage randomization analysis"
        workbook.properties.creator = "POCC-DeepMD manuscript analysis"
        workbook.properties.description = (
            "Source-backed results, draws, candidate frames, CIF index, and trajectory provenance."
        )
    temporary.replace(OUTPUT)
    receipt = {
        "schema_version": 1,
        "path": OUTPUT.name,
        "sha256": sha256(OUTPUT),
        "bytes": OUTPUT.stat().st_size,
        "sheets": ["README", "Figure", *tables.keys()],
        "row_counts": {name: len(dataframe) for name, dataframe in tables.items()},
    }
    (RESULTS / "workbook_receipt.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
